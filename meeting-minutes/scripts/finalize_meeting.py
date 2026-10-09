"""Validate internal CSV data and create a populated meeting.xlsx."""

from __future__ import annotations

import argparse
import csv
import shutil
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


REQUIRED_CSV_NAMES = ("header.csv", "minutes.csv", "participants.csv", "Agenda.csv")


def parse_args() -> argparse.Namespace:
    project_dir = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(
        description=(
            "work/finalの4種類のCSVを検証し、テンプレートへ値を直接書き込んだ"
            "output/meeting.xlsxを作成します。CSVはoutputへコピーしません。"
        )
    )
    parser.add_argument(
        "data_dir",
        nargs="?",
        type=Path,
        default=project_dir / "work" / "final",
        help="内部用の4 CSVがあるフォルダ（既定: work/final）",
    )
    parser.add_argument(
        "--template",
        type=Path,
        default=project_dir / "template" / "meeting.xlsx",
        help="書式の原本として使うExcelテンプレート",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=project_dir / "output" / "meeting.xlsx",
        help="完成Excelの保存先",
    )
    parser.add_argument("--force", action="store_true", help="既存の完成Excelを上書きする")
    return parser.parse_args()


def read_template_header(path: Path) -> list[str]:
    if not path.is_file():
        raise ValueError(f"CSVテンプレートが見つかりません: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        header = next(csv.reader(stream), None)
    if not header:
        raise ValueError(f"CSVテンプレートに列名がありません: {path}")
    return header


def read_and_validate_csv(path: Path, expected_header: list[str]) -> list[list[str]]:
    if not path.is_file():
        raise ValueError(f"必要なCSVが見つかりません: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        header = next(reader, None)
        if header != expected_header:
            raise ValueError(f"{path.name}の列名・列順がテンプレートと一致しません: {header}")
        rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        raise ValueError(f"{path.name}にデータ行がありません。")
    for number, row in enumerate(rows, start=2):
        if len(row) != len(expected_header):
            raise ValueError(
                f"{path.name}の{number}行目の列数が不正です: "
                f"期待={len(expected_header)}, 実際={len(row)}"
            )
    return rows


def clear_values(sheet, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in sheet.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None


def write_data_sheet(sheet, header: list[str], rows: list[list[str]]) -> None:
    clear_values(sheet, 1, max(sheet.max_row, len(rows) + 1), 1, max(sheet.max_column, len(header)))
    for column, value in enumerate(header, start=1):
        sheet.cell(row=1, column=column, value=value)
    for row_number, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_number, column=column, value=value)


def copy_row_style(sheet, source_row: int, target_row: int, max_column: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column in range(1, max_column + 1):
        source = sheet.cell(row=source_row, column=column)
        target = sheet.cell(row=target_row, column=column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)


def extend_agenda_area(sheet, agenda_count: int) -> int:
    extra_rows = max(0, agenda_count - 5)
    if extra_rows == 0:
        return 0

    shifted_merges = []
    for merged in list(sheet.merged_cells.ranges):
        if merged.min_row >= 15:
            shifted_merges.append(
                (
                    merged.min_row + extra_rows,
                    merged.min_col,
                    merged.max_row + extra_rows,
                    merged.max_col,
                )
            )
            sheet.unmerge_cells(str(merged))
    sheet.insert_rows(15, extra_rows)
    for min_row, min_col, max_row, max_col in shifted_merges:
        sheet.merge_cells(
            start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col
        )

    for row_number in range(14, 14 + extra_rows):
        copy_row_style(sheet, 13, row_number, 6)
        sheet.merge_cells(start_row=row_number, start_column=3, end_row=row_number, end_column=4)
        sheet.merge_cells(start_row=row_number, start_column=5, end_row=row_number, end_column=6)
    return extra_rows


def extend_minutes_area(sheet, minutes_count: int) -> None:
    extra_rows = max(0, minutes_count - 21)
    for offset in range(extra_rows):
        row_number = 25 + offset
        copy_row_style(sheet, 24, row_number, 3)


def populate_display_sheets(workbook, data: dict[str, list[list[str]]]) -> None:
    summary = workbook["議事録"]
    conversation = workbook["発言記録"]
    header = data["header.csv"][0]
    minutes = data["minutes.csv"]
    participants = data["participants.csv"]
    agenda = data["Agenda.csv"]

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = None

    summary["B3"] = header[0]
    summary["B4"] = f"{header[1]} {header[2]}～{header[3]}".strip()
    summary["E4"] = header[4]
    summary["B5"] = "、".join(row[2] for row in participants if row[2].strip())

    agenda_shift = extend_agenda_area(summary, len(agenda))
    extend_minutes_area(conversation, len(minutes))

    clear_values(summary, 9, 13 + agenda_shift, 2, 6)
    for row_number, row in enumerate(agenda, start=9):
        summary.cell(row=row_number, column=1, value=row_number - 8)
        summary.cell(row=row_number, column=2, value=row[0])
        summary.cell(row=row_number, column=3, value=row[1])
        summary.cell(row=row_number, column=5, value=row[2])

    action_start = 17 + agenda_shift
    clear_values(summary, action_start, action_start + 4, 2, 6)
    decisions = [row[2] for row in agenda if row[2].strip()]
    for row_number, decision in enumerate(decisions[:5], start=action_start):
        summary.cell(row=row_number, column=2, value=decision)

    clear_values(conversation, 4, max(24, len(minutes) + 3), 2, 3)
    for row_number, row in enumerate(minutes, start=4):
        conversation.cell(row=row_number, column=1, value=row_number - 3)
        conversation.cell(row=row_number, column=2, value=row[4])
        conversation.cell(row=row_number, column=3, value=row[5])


def main() -> int:
    args = parse_args()
    data_dir = args.data_dir.resolve()
    template = args.template.resolve()
    target = args.output.resolve()

    try:
        if not template.is_file():
            raise ValueError(f"Excelテンプレートが見つかりません: {template}")
        if target.exists() and not args.force:
            raise FileExistsError(
                f"完成Excelは既に存在します: {target}\n"
                "上書きする場合は--forceを指定してください。"
            )

        template_dir = template.parent
        headers: dict[str, list[str]] = {}
        data: dict[str, list[list[str]]] = {}
        for name in REQUIRED_CSV_NAMES:
            headers[name] = read_template_header(template_dir / name)
            data[name] = read_and_validate_csv(data_dir / name, headers[name])

        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(template, target)
        workbook = load_workbook(target)

        if "_config" in workbook.sheetnames:
            workbook.remove(workbook["_config"])
        for defined_name in list(workbook.defined_names):
            del workbook.defined_names[defined_name]

        write_data_sheet(workbook["header"], headers["header.csv"], data["header.csv"])
        write_data_sheet(workbook["minutes"], headers["minutes.csv"], data["minutes.csv"])
        write_data_sheet(
            workbook["participants"], headers["participants.csv"], data["participants.csv"]
        )
        write_data_sheet(workbook["Agenda"], headers["Agenda.csv"], data["Agenda.csv"])
        populate_display_sheets(workbook, data)

        for sheet_name in ("header", "minutes", "participants", "Agenda"):
            workbook[sheet_name].sheet_state = "hidden"
        workbook.calculation.fullCalcOnLoad = False
        workbook.calculation.forceFullCalc = False
        workbook.save(target)

        print(f"完成Excel: {target}")
        for name in REQUIRED_CSV_NAMES:
            print(f"{name}: {len(data[name])}件（内部データ、outputへのコピーなし）")
    except (FileExistsError, KeyError, OSError, ValueError) as exc:
        print(f"エラー: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
